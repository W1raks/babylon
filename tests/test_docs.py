"""Перевод документов: оформление остаётся, текст меняется (перевод подменён)."""
import docx
import pymupdf as fitz
import pytest

from babylon import docs

TR = {
    "Report on languages": "Отчёт о языках",
    "Most people give up after a few weeks.": "Большинство бросает через несколько недель.",
    "Language": "Язык",
    "Hours": "Часы",
    "English": "Английский",
    "The best way to learn a language is to use it every day, even for ten minutes.":
        "Лучший способ выучить язык — пользоваться им каждый день, хотя бы десять минут.",
    "Inside a blue box": "Внутри синей рамки",
    "Hello there!": "Привет!",
    "How are you today?": "Как ты сегодня?",
}


def fake_many(texts):
    return [TR.get(t, "") for t in texts]


def test_docx_keeps_styles_and_tables(tmp_path):
    d = docx.Document()
    d.add_heading("Report on languages", level=1)
    p = d.add_paragraph()
    r1 = p.add_run("Most people give ")
    r1.bold = True
    p.add_run("up after a few weeks.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Language", "Hours"
    t.cell(1, 0).text, t.cell(1, 1).text = "English", "120"
    src, out = tmp_path / "a.docx", tmp_path / "b.docx"
    d.save(src)
    n = docs.translate_file(str(src), str(out), fake_many)
    assert n == 5
    r = docx.Document(out)
    assert r.paragraphs[0].text == "Отчёт о языках" and r.paragraphs[0].style.name == "Heading 1"
    assert r.paragraphs[1].text == "Большинство бросает через несколько недель."
    assert r.tables[0].cell(0, 0).text == "Язык" and r.tables[0].cell(1, 1).text == "120"


def test_pdf_replaces_text_keeps_drawing(tmp_path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox(fitz.Rect(72, 72, 400, 140),
                        "The best way to learn a language is to use it every day, even for ten minutes.", fontsize=13)
    page.draw_rect(fitz.Rect(60, 200, 400, 260), color=(0, 0, 1), fill=(0.9, 0.9, 1))
    page.insert_text((72, 235), "Inside a blue box", fontsize=18, color=(0.6, 0, 0))
    src, out = tmp_path / "a.pdf", tmp_path / "b.pdf"
    doc.save(src)
    n = docs.translate_file(str(src), str(out), fake_many)
    assert n == 2
    text = " ".join(fitz.open(out)[0].get_text().split())
    assert "Лучший способ выучить язык" in text and "Внутри синей рамки" in text
    assert "The best way" not in text
    drawings = fitz.open(out)[0].get_drawings()
    assert any(abs(dr["rect"].x0 - 60) < 1 for dr in drawings)  # синяя рамка на месте


def test_pdf_scan_is_reported(tmp_path):
    doc = fitz.open()
    doc.new_page().draw_rect(fitz.Rect(10, 10, 50, 50), fill=(0, 0, 0))
    doc.save(tmp_path / "scan.pdf")
    with pytest.raises(docs.DocError):
        docs.translate_file(str(tmp_path / "scan.pdf"), str(tmp_path / "o.pdf"), fake_many)


def test_srt_and_txt(tmp_path):
    (tmp_path / "a.srt").write_text("1\n00:00:01,000 --> 00:00:02,500\nHello there!\n\n"
                                    "2\n00:00:03,000 --> 00:00:04,000\n<i>How are you</i>\ntoday?\n", encoding="utf-8")
    assert docs.translate_file(str(tmp_path / "a.srt"), str(tmp_path / "b.srt"), fake_many) == 2
    assert (tmp_path / "b.srt").read_text(encoding="utf-8") == (
        "1\n00:00:01,000 --> 00:00:02,500\nПривет!\n\n2\n00:00:03,000 --> 00:00:04,000\nКак ты сегодня?\n\n")
    (tmp_path / "a.txt").write_text("Hello there!\n\nHow are you\ntoday?\n", encoding="utf-8")
    docs.translate_file(str(tmp_path / "a.txt"), str(tmp_path / "b.txt"), fake_many)
    assert (tmp_path / "b.txt").read_text(encoding="utf-8") == "Привет!\n\nКак ты сегодня?\n"


def test_out_name():
    assert docs.out_name("C:/x/report.docx", "ru") == "C:/x/report (перевод ru).docx"
